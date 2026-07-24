"""Paso 5-6: generar el reporte de impugnacion (Renta Fija.xlsx).

>>> ESTE ES EL UNICO PASO QUE QUEDA PENDIENTE DE DESARROLLO <<<
Ver PENDIENTE_PASO_5.md en esta misma carpeta.

La macro VBA original (`Informacion Impugnacion.xlsm`) lee el archivo plano
descargado de INFOVALMER (SXMMDDYY) y arma `Renta Fija.xlsx`. La decision
acordada es MIGRAR esa logica a Python (pandas/openpyxl), no ejecutar el VBA.

Como todavia no se ha inspeccionado el .xlsm, aqui hay dos implementaciones:

  1. TransformadorImpugnacionPython  -> la real. Hoy lanza PasoPendienteError
     porque falta traducir la logica del VBA. Este es el trabajo restante.

  2. TransformadorPlaceholder        -> genera un Renta Fija.xlsx MINIMO y
     claramente marcado como placeholder, para poder correr y validar el resto
     del pipeline (descarga -> INFOVALMER -> validaciones -> correo) end-to-end
     en modo test sin quedar bloqueados por el paso 5.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from core.adapters.excel_base import TransformadorImpugnacion
from core.contract import PasoPendienteError


class TransformadorImpugnacionPython(TransformadorImpugnacion):
    """Reimplementacion en Python de la macro. PENDIENTE (paso 5)."""

    def __init__(self, archivo_salida: str = "Renta Fija.xlsx", logger=None) -> None:
        self.archivo_salida = archivo_salida
        self.logger = logger

    def generar_reporte(self, archivo_entrada: Path, carpeta_salida: Path) -> Path:
        raise PasoPendienteError(
            "Paso 5 pendiente: falta migrar la logica de la macro "
            "'Informacion Impugnacion.xlsm' a Python. Ver PENDIENTE_PASO_5.md. "
            "Para probar el resto del pipeline usar TransformadorPlaceholder "
            "(config: excel.motor='placeholder')."
        )


class TransformadorPlaceholder(TransformadorImpugnacion):
    """Genera un Renta Fija.xlsx minimo marcado como PLACEHOLDER.

    NO calcula la impugnacion real: solo permite ejercitar validaciones y envio.
    """

    def __init__(self, archivo_salida: str = "Renta Fija.xlsx", logger=None) -> None:
        self.archivo_salida = archivo_salida
        self.logger = logger

    def generar_reporte(self, archivo_entrada: Path, carpeta_salida: Path) -> Path:
        from openpyxl import Workbook

        archivo_entrada = Path(archivo_entrada)
        carpeta_salida = Path(carpeta_salida)
        carpeta_salida.mkdir(parents=True, exist_ok=True)
        destino = carpeta_salida / self.archivo_salida

        wb = Workbook()
        ws = wb.active
        ws.title = "Impugnacion"
        ws["A1"] = "PLACEHOLDER - PASO 5 PENDIENTE (macro por migrar a Python)"
        ws["A2"] = f"Generado: {datetime.now():%Y-%m-%d %H:%M:%S}"
        ws["A3"] = f"Archivo de entrada: {archivo_entrada.name}"
        ws["A4"] = "Este archivo NO contiene la impugnacion real."
        wb.save(destino)

        if self.logger:
            self.logger.warning(
                "transform_placeholder",
                extra={"detalle": f"Renta Fija.xlsx PLACEHOLDER generado en {destino}. "
                                  "Paso 5 pendiente."},
            )
        return destino


def crear_transformador(
    motor: str, archivo_salida: str, logger=None, cfg_excel=None
) -> TransformadorImpugnacion:
    """Fabrica segun config.excel.motor: 'placeholder' | 'python' | 'com'."""
    if motor == "placeholder":
        return TransformadorPlaceholder(archivo_salida, logger)
    if motor == "com":
        # Fallback Windows: correr el VBA original. Requiere nombre_macro confirmado.
        from pathlib import Path

        from core.adapters.excel_com import TransformadorViaCOM
        nombre_macro = getattr(cfg_excel, "nombre_macro", "<PENDIENTE_CONFIRMAR>")
        ruta_xlsm = Path(getattr(cfg_excel, "archivo_macro", "Informacion Impugnacion.xlsm"))
        return TransformadorViaCOM(ruta_xlsm, nombre_macro, archivo_salida,
                                   getattr(cfg_excel, "timeout_seg", 300), logger)
    return TransformadorImpugnacionPython(archivo_salida, logger)
